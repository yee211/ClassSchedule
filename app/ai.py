"""通过 OpenAI 兼容协议调用大模型，把 Excel 课表结构化为固定 schema。

换厂商只需改 .env 里的 AI_BASE_URL / AI_API_KEY / AI_MODEL 三项，代码无需调整。
AI_API_KEY 为空时本模块直接返回 ai-not-configured，调用方会自动回退到
app/parser.py 的确定性解析，因此项目在未配置任何密钥时依然完整可用。
"""
import json
import os
import re
from pathlib import Path

import httpx
from dotenv import load_dotenv
from openpyxl.utils import get_column_letter
from pydantic import BaseModel, Field, ValidationError, field_validator, model_validator

from .excel import MAX_ROWS_PER_SHEET, read_sheets
from .parser import course_color, merge_section_courses, parse_weeks


ROOT = Path(__file__).resolve().parent.parent
load_dotenv(ROOT / ".env")

AI_BASE_URL = os.getenv("AI_BASE_URL", "https://api.deepseek.com/v1").rstrip("/")
AI_API_KEY = os.getenv("AI_API_KEY", "").strip().strip('"').strip("'")
AI_MODEL = os.getenv("AI_MODEL", "deepseek-chat")

AI_TIMEOUT = float(os.getenv("AI_TIMEOUT_SECONDS", "120"))
AI_MAX_INPUT_CHARS = int(os.getenv("AI_MAX_INPUT_CHARS", "60000"))

# 单个 sheet 的行数上限，防止带大量尾部格式的空表拖垮序列化
MAX_ROWS_PER_SHEET = 1500

SYSTEM_PROMPT = """你是教务课表结构化解析器。用户会提供一个 Excel 工作簿的单元格矩阵文本，格式为逐 sheet 的
`=== SHEET n: "表名" (rows=.., cols=..) ===`，其下可能有 `MERGED:` 行列出合并单元格范围，
再往下每行形如 `ROW <行号>: <坐标>="<值>" <坐标>="<值>" ...`，值内的 `\\n` 表示原单元格中的换行。

你的任务是从中提取课程安排，并且只输出一个 JSON 对象，禁止输出 markdown 代码块、注释或任何解释文字。

输出 JSON 结构：
{
  "name": "课表名称，如“张三的课表”；无法判断时填“我的课表”",
  "term": "学期标识，优先取标题中的“YYYY-YYYY学年第N学期”；确实无法判断时填“导入课表”",
  "courses": [
    {
      "name": "课程名称，必填",
      "teacher": "授课教师，缺失填空字符串",
      "room": "教室或上课地点，缺失填空字符串",
      "weekday": 1,
      "start_section": 1,
      "end_section": 2,
      "weeks": "2-16(双)"
    }
  ]
}

解析规则：
1. weekday：1=周一、2=周二 …… 7=周日。“星期一/周一/礼拜一”等写法一律映射为对应数字。
2. start_section / end_section：1 到 12 的整数，取自“节次”列或行标签（如“第1-2节”“1-2”“上午第1节”）。
3. 连堂必须合并成一条记录：同一课程、同一教师、同一教室且节次连续（如 1-2 节与 3-4 节）时，
   输出 start_section=1、end_section=4，不要拆成两条。
4. weeks：**原样抄录表中的周次表达式字符串，不要展开成数字数组**。例如 "1-16"、"2-16(双)"、
   "3-17(单)"、"1,3,5-8"、"11-15"、"1"。保留单/双标记与括号，可去掉“第”“周”等冗余字。
   若确实无法判断周次，输出空字符串 ""。
5. 同一个单元格里含多门课程时（如同一时段上下两段不同课程），拆成多条记录。
6. teacher / room 缺失时输出空字符串，严禁编造，也不要从相邻行推断。
7. 无法确定 weekday 或节次的行直接丢弃，不要猜测。
8. 忽略表头行、序号列、说明、备注、合计、统计、签名、审核等非课程信息。
9. 课程名称要清理多余空白与换行，但不要把课程名与教师名粘连在一起。
10. 若整份文件里确实没有任何课程，courses 输出空数组。"""


class AICourse(BaseModel):
    """与 courses 表字段严格对齐；color 由后端确定性生成，不向模型索要。"""
    name: str
    teacher: str = ""
    room: str = ""
    weekday: int = Field(ge=1, le=7)
    start_section: int = Field(ge=1, le=12)
    end_section: int = Field(ge=1, le=12)
    weeks: list[int] = Field(default_factory=list)

    @field_validator("name", "teacher", "room")
    @classmethod
    def clean_text(cls, value, info):
        # 截断到数据库列宽而不是直接报错，避免个别超长字段拖垮整批导入
        limit = {"name": 80, "teacher": 40, "room": 40}[info.field_name]
        return re.sub(r"\s+", " ", str(value or "")).strip()[:limit]

    @field_validator("weeks", mode="before")
    @classmethod
    def coerce_weeks(cls, value):
        # 提示词要求模型直接回传 "1-16" / "2-16(双)" 这样的原式（输出 token 比展开数组少一个量级，
        # 直接决定了接口耗时），用本地解析器展开；兼容模型仍然返数组的情况
        if isinstance(value, str):
            return parse_weeks(value)
        return value

    @field_validator("weeks")
    @classmethod
    def clean_weeks(cls, value):
        return sorted({week for week in value if isinstance(week, int) and 1 <= week <= 30})

    @model_validator(mode="after")
    def order_sections(self):
        # 数据库有 CHECK (end_section >= start_section)，模型偶发倒序时在此纠正
        if self.end_section < self.start_section:
            self.start_section, self.end_section = self.end_section, self.start_section
        return self


def serialize_workbook(path: Path) -> str:
    """Render every sheet as coordinate-tagged text rows for the model to read."""
    sheets = read_sheets(path)
    chunks = []
    used = 0
    truncated = False
    for index, sheet in enumerate(sheets, 1):
        lines = [f'=== SHEET {index}: "{sheet.title}" (rows={sheet.max_row}, cols={sheet.max_column}) ===']
        if sheet.merged:
            lines.append("MERGED: " + " ".join(sheet.merged))
        for row_index, cells in sheet.rows:
            if row_index > MAX_ROWS_PER_SHEET:
                break
            cell_texts = []
            for col_index, val in sorted(cells.items()):
                text = str(val or "").strip()
                if not text:
                    continue
                # 换行转义为字面 \n，既保留“教师/周次分行”信息又不破坏 ROW 单行结构
                text = text.replace("\r\n", "\n").replace("\r", "\n").replace("\n", "\\n")
                coord = f"{get_column_letter(col_index)}{row_index}"
                cell_texts.append(f'{coord}="{text}"')
            if not cell_texts:
                continue
            lines.append(f"ROW {row_index}: " + " ".join(cell_texts))
        block = "\n".join(lines)
        if used + len(block) > AI_MAX_INPUT_CHARS:
            truncated = True
            break
        chunks.append(block)
        used += len(block) + 1
    if truncated:
        chunks.append("... [TRUNCATED]")
    return "\n".join(chunks)



def _extract_json(text: str) -> str:
    """Defensively unwrap code fences or surrounding prose before json.loads."""
    text = (text or "").strip()
    fence = re.match(r"^```(?:json)?\s*(.*?)\s*```$", text, re.DOTALL)
    if fence:
        text = fence.group(1).strip()
    start, end = text.find("{"), text.rfind("}")
    return text[start:end + 1] if start != -1 and end > start else text


def _coerce_schedule(payload) -> dict | None:
    """Validate model output course by course; skip bad rows instead of failing all."""
    if not isinstance(payload, dict) or not isinstance(payload.get("courses"), list):
        return None
    courses = []
    for item in payload["courses"]:
        if not isinstance(item, dict):
            continue
        try:
            course = AICourse.model_validate(item)
        except ValidationError:
            continue
        if not course.name:
            continue
        data = course.model_dump()
        data["color"] = course_color(data["weekday"], data["start_section"])
        courses.append(data)
    # term 是 (user_id, term) 幂等覆盖键，必须稳定，不能落到随机文件名上
    term = re.sub(r"\s+", " ", str(payload.get("term") or "")).strip()[:80] or "导入课表"
    name = re.sub(r"\s+", " ", str(payload.get("name") or "")).strip()[:80] or "我的课表"
    return {"name": name, "term": term, "courses": merge_section_courses(courses)}


def parse_with_ai(path: Path) -> tuple[dict | None, str]:
    """Parse an Excel timetable through an OpenAI-compatible chat model.

    Never raises: every failure mode is reported through the engine tag so the
    caller can silently fall back to the deterministic parser.
    """
    if not AI_API_KEY:
        return None, "ai-not-configured"
    try:
        serialized = serialize_workbook(path)
    except Exception:
        return None, "ai-unreadable"
    if not serialized.strip():
        return None, "ai-empty-workbook"
    try:
        response = httpx.post(
            f"{AI_BASE_URL}/chat/completions",
            headers={"Authorization": f"Bearer {AI_API_KEY}"},
            timeout=AI_TIMEOUT,
            json={
                "model": AI_MODEL,
                "temperature": 0,
                "response_format": {"type": "json_object"},
                "messages": [
                    {"role": "system", "content": SYSTEM_PROMPT},
                    {"role": "user", "content": serialized},
                ],
            },
        )
    except Exception:
        # 契约是“永不抛出”，因此不只捕 httpx.HTTPError：非法 URL、不支持的协议、
        # 重定向与解码异常等也应降级为兜底解析，而不是变成 500
        return None, "ai-unavailable"
    if response.status_code != 200:
        return None, f"ai-http-{response.status_code}"
    try:
        content = response.json()["choices"][0]["message"]["content"]
    except (ValueError, KeyError, IndexError, TypeError):
        return None, "ai-invalid-response"
    try:
        payload = json.loads(_extract_json(content))
    except ValueError:
        return None, "ai-invalid-json"
    schedule = _coerce_schedule(payload)
    if not schedule:
        return None, "ai-invalid-schema"
    if not schedule["courses"]:
        return None, "ai-empty"
    return schedule, "ai"
